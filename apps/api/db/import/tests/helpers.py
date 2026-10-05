"""Shared test helpers: synthetic grids, workbooks and seed inputs."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from clean_products import is_credit
from seed_model import Inputs
from sheets import Grid

Doc = dict[str, Any]


def grid(cells: dict[str, object], rows: int | None = None, cols: int = 20) -> Grid:
    """Build a 0-based grid from {"D3": value} cells."""
    coords = {coordinate_from_string(k): v for k, v in cells.items()}
    height = rows or max((r for (_, r) in coords), default=1)
    g: Grid = [[None] * cols for _ in range(height)]
    for (col, row), v in coords.items():
        g[row - 1][column_index_from_string(col) - 1] = v
    return g


def save_workbook(path: Path, sheets: dict[str, dict[str, object]]) -> Path:
    """Write a workbook with one sheet per {name: cells} entry."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for k, v in cells.items():
            ws[k] = v
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# Seed inputs
def line(no: int, sell: float | None, cost: float | None = None, **kw: object) -> Doc:
    base: Doc = {
        "line_no": no,
        "row": 14 + no,
        "qty": 1.0,
        "unit": "PCS",
        "request": f"Item {no}",
        "offer": None,
        "impa": None,
        "sell": sell,
        "cost": cost,
        "line_total": sell,
        "in_total": True,
        "available": True,
        "vendor": None,
        "channel": None,
    }
    return base | kw


def staged(
    rid: str,
    number: str,
    day: str | None,
    lines: list[Doc],
    client: str = "PT. Alpha",
    **kw: object,
) -> Doc:
    rec: Doc = {
        "id": rid,
        "source": {"files": [rid.split("#")[0]]},
        "number": {"original": number, "print": number},
        "year": int((day or "2025")[:4]),
        "date": day,
        "client": {"name": client, "broker_note": None},
        "contact": {"name": "Ibu Sari", "email": "sari@alpha.example", "phone": None},
        "client_ref": None,
        "vessel": None,
        "terms": {
            "payment": "30 days",
            "validity_days": 3,
            "delivery_time": None,
            "delivery_place": None,
        },
        "discount": {"kind": "none", "pct": None, "amount": 0},
        "lines": lines,
        "extras": [],
        "notes": [],
        "revision": None,
        "print_totals": None,
        "pdf": None,
        "computed_totals": {"ppn_rate": 11.0},
        "reconciliation": [],
        "aliases": [],
        "flags": [],
    }
    return rec | kw


def inputs(
    records: list[Doc], pos: list[Doc] | None = None, invoices: list[Doc] | None = None
) -> Inputs:
    products = [
        {
            "id": f"P{i:04d}",
            "name": ln["request"],
            "description": None,
            "pack_note": None,
            "impa_code": None,
            "kind": "product",
            "default_unit": "PCS",
            "source_lines": [{"quotation": r["id"], "line_no": ln["line_no"]}],
        }
        for i, (r, ln) in enumerate(
            ((r, ln) for r in records for ln in r["lines"] if not is_credit(ln)), 1
        )
    ]
    return Inputs(records, products, {}, pos or [], invoices or [])


def po_doc(rec: Doc, lines: list[Doc], **kw: object) -> Doc:
    """A client PO ordering `lines` ({no, qty, unit_price, quote_row}) of `rec`."""
    po: Doc = {
        "id": f"po:{rec['id']}",
        "client": rec["client"]["name"],
        "client_printed": rec["client"]["name"].upper(),
        "client_po_number": "PO-1",
        "client_po_number_missing_reason": None,
        "po_date": rec["date"],
        "user_end": None,
        "delivery_place": None,
        "payment_terms": "30 days",
        "client_reference": None,
        "charges": [],
        "discount": 0,
        "untaxed": 0,
        "source_files": [{"path": "PO/po.pdf", "part": None}],
        "synthetic_quotation": None,
        "notes": [],
        "link": {"kind": "single"},
        "quotation": {
            "file": "Quotation/" + rec["source"]["files"][0],
            "printed_number": rec["number"]["original"],
            "lines": [{"po_line": ln["no"], "quote_row": ln["quote_row"]} for ln in lines],
            "charges": [],
            "other_files": [],
        },
        "lines": [
            {
                "no": ln["no"],
                "description": ln.get("description", f"Item {ln['no']}"),
                "qty": ln["qty"],
                "unit": ln.get("unit", "PCS"),
                "unit_price": ln["unit_price"],
            }
            for ln in lines
        ],
    }
    return po | kw


def invoice_doc(po: Doc, number: str, day: str, **kw: object) -> Doc:
    """The untaxed-free invoice billing every priced line of `po` at PPN 11%."""
    lines = [
        {
            "no": i,
            "description": ln["description"],
            "qty": ln["qty"],
            "unit": ln["unit"],
            "unit_price": ln["unit_price"],
            "amount": ln["qty"] * ln["unit_price"],
        }
        for i, ln in enumerate((ln for ln in po["lines"] if ln["unit_price"]), 1)
    ]
    dpp = sum(ln["amount"] for ln in lines) - po["discount"]
    ppn = round(dpp * 0.11, 2)
    inv: Doc = {
        "invoice_number": number,
        "do_number": number.replace("INV", "DO"),
        "invoice_date": day,
        "due_date": None,
        "do_date": day,
        "extra_dos": [],
        "po_id": po["id"],
        "buyer": {"name": po["client_printed"], "address": None, "npwp": None},
        "lines": lines,
        "untaxed": 0,
        "discount": po["discount"],
        "dpp": dpp,
        "ppn_rate": 11,
        "ppn": ppn,
        "total": dpp + ppn,
        "source_file": f"Invoice/{number}.xlsx",
    }
    return inv | kw
