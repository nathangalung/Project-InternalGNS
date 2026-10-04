"""Unit tests for building quotation records."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest

from helpers import save_workbook
from quotation import (
    _apply_print_line_set,
    _apply_print_no_offer,
    _tax_setup,
    build_workbook,
    compute_totals,
    normalize_line,
    number_flags,
    parse_days,
    resolve_discount,
    split_vessel_lead,
    vessel_from_text,
)
from sheets import Line, PrintInfo, PrintTotals


def _line(**kw: object) -> Line:
    base: dict[str, object] = {
        "row": 14,
        "no": "1",
        "qty": 2.0,
        "unit": "Pcs",
        "request": "Bolt M10",
        "offer": None,
    }
    base.update(kw)
    return Line(**base)  # type: ignore[arg-type]


def test_normalize_line_basic() -> None:
    ln = normalize_line(
        _line(impa=691185, sell=5000.0, sell_amount=10000.0, cost=3000.0, vendor="Toko A"),
        has_amount=True,
    )
    assert ln["request"] == "Bolt M10"
    assert ln["offer"] is None
    assert (ln["impa"], ln["impa_problem"], ln["unit"], ln["unit_mapped"]) == (
        "691185",
        None,
        "PCS",
        True,
    )
    assert (ln["sell"], ln["cost"], ln["line_total"], ln["in_total"]) == (
        5000.0,
        3000.0,
        10000.0,
        True,
    )
    assert ln["available"] is True
    assert ln["vendor"] == "Toko A"
    assert ln["flags"] == []


def test_normalize_line_no_offer() -> None:
    ln = normalize_line(
        _line(offer="No offer. Prescription required", sell=0.0, sell_amount=0.0),
        has_amount=True,
    )
    assert ln["available"] is False
    assert ln["offer"] is None
    assert ln["offer_note"] == "Prescription required"
    assert (ln["sell"], ln["line_total"], ln["in_total"]) == (0.0, 0.0, True)


def test_normalize_line_no_offer_in_nama_asli_only() -> None:
    ln = normalize_line(_line(nama_asli="No Offer", sell=0.0), has_amount=True)
    assert ln["available"] is False


def test_normalize_line_offer_only_layout_uses_offer_as_request() -> None:
    ln = normalize_line(_line(request=None, offer="Cylinder Head Yanmar"), has_amount=False)
    assert (ln["request"], ln["offer"]) == ("Cylinder Head Yanmar", "Cylinder Head Yanmar")


def test_normalize_line_impa_checks() -> None:
    ln = normalize_line(
        _line(impa="#REF!", offer="Fan IMPA 591418, 220V", request="Blower impa 591506"),
        has_amount=False,
    )
    assert (ln["impa"], ln["impa_raw"], ln["impa_problem"]) == (None, "#REF!", "invalid")
    assert ln["impa_in_text"] == ["591506", "591418"]
    assert "impa_invalid" in ln["flags"]
    ln = normalize_line(_line(impa=591506, offer="Impa 591418 fan"), has_amount=False)
    assert "impa_text_differs" in ln["flags"]


def test_normalize_line_amount_wins_over_qty_times_price() -> None:
    ln = normalize_line(_line(qty=100.0, sell=36000.0, sell_amount=360000.0), has_amount=True)
    assert (ln["sell"], ln["sell_printed"], ln["line_total"]) == (3600.0, 36000.0, 360000.0)
    assert "unit_price_from_amount" in ln["flags"]


def test_normalize_line_amount_and_cost_amount_correct_the_qty() -> None:
    # Q-528 row 19: 3 requested, 1 available; both amounts are for one unit.
    ln = normalize_line(
        _line(qty=3.0, sell=5400000.0, sell_amount=5400000.0, cost=3000000.0,
              cost_amount=3000000.0),
        has_amount=True,
    )  # fmt: skip
    assert (ln["qty"], ln["qty_raw"], ln["sell"], ln["cost"]) == (1.0, 3.0, 5400000.0, 3000000.0)
    assert "qty_from_amount" in ln["flags"]


def test_normalize_line_rescaled_price_keeps_a_consistent_cost() -> None:
    ln = normalize_line(
        _line(qty=100.0, sell=36000.0, sell_amount=360000.0, cost=1500.0, cost_amount=150000.0),
        has_amount=True,
    )
    assert (ln["qty"], ln["sell"], ln["cost"]) == (100.0, 3600.0, 1500.0)
    ln = normalize_line(
        _line(qty=100.0, sell=36000.0, sell_amount=360000.0, cost=15000.0), has_amount=True
    )
    assert (ln["sell"], ln["cost"]) == (3600.0, 1500.0)
    assert "cost_rescaled_with_price" in ln["flags"]


def test_normalize_line_cost_from_cost_amount() -> None:
    ln = normalize_line(
        _line(qty=2.0, sell=250000.0, sell_amount=500000.0, cost=None, cost_amount=380000.0),
        has_amount=True,
    )
    assert ln["cost"] == 190000.0 and "cost_from_amount" in ln["flags"]


def test_normalize_line_without_amount_is_not_in_total() -> None:
    ln = normalize_line(_line(no=None, qty=1.0, sell=30000.0, sell_amount=None), has_amount=True)
    assert ln["in_total"] is False
    assert ln["line_total"] == 0.0
    assert "not_in_amount_column" in ln["flags"]


def test_normalize_line_negative_amount() -> None:
    ln = normalize_line(
        _line(no=None, qty=None, request=None, offer="Trade-in", sell=None, sell_amount=-500.0),
        has_amount=True,
    )
    assert (ln["qty"], ln["sell"], ln["line_total"]) == (1.0, -500.0, -500.0)
    assert "negative_amount" in ln["flags"]


def test_normalize_line_unit_fallback_and_qty_flags() -> None:
    ln = normalize_line(_line(unit="Widget", qty=0.0), has_amount=False)
    assert (ln["unit"], ln["unit_mapped"]) == ("OTH", False)
    assert "unit_unmapped" in ln["flags"] and "qty_zero" in ln["flags"]


@pytest.mark.parametrize(
    ("label", "pct_label", "amount", "gross", "kind", "pct", "value", "flag"),
    [
        ("Diskon 5%", 5.0, 50.0, 1000.0, "pct", 5.0, 50.0, None),
        ("Diskon 5%", 5.0, 0.0, 1000.0, "none", None, 0.0, "discount_label_without_amount"),
        ("Diskon", None, 100.0, 1000.0, "pct", 10.0, 100.0, "discount_pct_inferred"),
        ("Diskon", None, 1240000.0, 30170000.0, "amount", None, 1240000.0, None),
        ("Diskon 5%", 5.0, 70.0, 1000.0, "amount", None, 70.0, "discount_label_pct_mismatch"),
        ("Diskon", None, 0.0, 1000.0, "none", None, 0.0, None),
    ],
)
def test_resolve_discount(
    label: str,
    pct_label: float | None,
    amount: float,
    gross: float,
    kind: str,
    pct: float | None,
    value: float,
    flag: str | None,
) -> None:
    t = PrintTotals(
        gross=gross, discount=amount, discount_label=label, discount_pct_label=pct_label
    )
    flags: list[str] = []
    d = resolve_discount(t, gross, flags)
    assert (d["kind"], d["pct"], d["amount"]) == (kind, pct, value)
    assert flags == ([flag] if flag else [])


def test_resolve_discount_without_print() -> None:
    flags: list[str] = []
    assert resolve_discount(None, 1000.0, flags)["kind"] == "none"


def test_compute_totals_2025_dpp() -> None:
    lines = [{"line_total": 9892000.0, "in_total": True}]
    t = compute_totals(lines, {"kind": "none", "pct": None, "amount": 0.0}, 12.0, True, [])
    assert t["gross"] == 9892000.0
    assert t["net"] == 9892000.0
    assert t["dpp"] == pytest.approx(9067666.6667, abs=0.01)
    assert t["ppn"] == pytest.approx(1088120.0, abs=0.01)
    assert t["grand"] == pytest.approx(10980120.0, abs=0.01)


def test_compute_totals_2024_discount_and_delivery() -> None:
    lines = [
        {"line_total": 6000000.0, "in_total": True},
        {"line_total": 99.0, "in_total": False},
    ]
    disc = {"kind": "pct", "pct": 5.0, "amount": 300000.0}
    t = compute_totals(lines, disc, 11.0, False, [("Delivery", 1000000.0, False)])
    assert (t["gross"], t["discount"], t["net"], t["dpp"]) == (6000000.0, 300000.0, 5700000.0, None)
    assert t["ppn"] == pytest.approx(627000.0)
    assert t["grand"] == pytest.approx(7327000.0)


@pytest.mark.parametrize(
    ("text", "want"),
    [
        ("MV YUXIN SATU", "MV YUXIN SATU"),
        ("SHIP: MV YUXIN SATU", "MV YUXIN SATU"),
        ("Untuk Kapal DAIDAN PERTIWI", "DAIDAN PERTIWI"),
        ("Untuk MV Daidan Mustikawati", "MV Daidan Mustikawati"),
        ("For MV: Dewi Saraswati", "MV Dewi Saraswati"),
        ("Untuk: Vessel Dewi Saraswati", "Dewi Saraswati"),
        ("TB Marina 18/BG BEE POWER 3005", "TB Marina 18/BG BEE POWER 3005"),
        ("SHIP: MV", None),
        ("Note: Material NBR", None),
        ("Hand Tools & Consumable Engine", None),
        ("form 305", None),
        ("TB Repair", None),
    ],
)
def test_vessel_from_text(text: str, want: str | None) -> None:
    assert vessel_from_text(text) == want


def test_split_vessel_lead() -> None:
    text = "MV YUXIN SATU\n\nHAMMER CHIPPING HANDLED 453GRM, NON-SPARK"
    assert split_vessel_lead(text) == ("MV YUXIN SATU", "HAMMER CHIPPING HANDLED 453GRM, NON-SPARK")
    assert split_vessel_lead("Hammer\nMV spec") == (None, "Hammer\nMV spec")
    assert split_vessel_lead("MV YUXIN SATU") == (None, "MV YUXIN SATU")
    assert split_vessel_lead(None) == (None, None)


def test_vessel_from_text_known_names() -> None:
    assert vessel_from_text("Marina 18", known={"MARINA 18"}) == "Marina 18"
    assert vessel_from_text("Stock Office", known={"MARINA 18"}) is None


@pytest.mark.parametrize(
    ("printed", "when", "file_number", "folder_year", "flags"),
    [
        ("Q-509/GNS/VII/2024", date(2024, 7, 26), "509", 2024, []),
        ("Q-963173-O/GNS/IV/2026", date(2026, 4, 30), "963173", 2026, []),
        ("Q-963046-O/6GNS/I/2026", date(2026, 1, 5), "963046", 2026, ["number_malformed"]),
        ("Q-758/MH/IX/2023", date(2024, 9, 2), "548", 2024, [
            "number_malformed", "number_year_ne_date", "file_number_ne_printed",
        ]),
        ("Q-951287/GNS/XI/2025", date(2025, 9, 20), "951287", 2025, ["number_month_ne_date"]),
        ("Q-95393/GNS/III/2025", date(2025, 3, 3), "953393", 2025, [
            "number_digits_unusual", "file_number_ne_printed",
        ]),
        ("Q-961035/GNS/I/2026", date(2025, 1, 20), "965034", 2026, [
            "number_year_ne_date", "date_year_ne_folder", "file_number_ne_printed",
        ]),
        (None, None, "546", 2024, ["number_missing", "date_missing"]),
    ],
)  # fmt: skip
def test_number_flags(
    printed: str | None,
    when: date | None,
    file_number: str | None,
    folder_year: int,
    flags: list[str],
) -> None:
    assert number_flags(printed, when, file_number, folder_year) == flags


@pytest.mark.parametrize(
    ("text", "want"),
    [("3 days", 3), ("30 Days", 30), ("3 DAYS", 3), ("COD", None), (None, None)],
)
def test_parse_days(text: str | None, want: int | None) -> None:
    assert parse_days(text) == want


def _workbook(tmp_path: Path) -> Path:
    data = {
        "A2": "Customer",
        "D2": "PT Kemala Shipping/MBSS",
        "A3": "No",
        "D3": "Q-968208-O/GNS/VI/2026",
        "A4": "Date",
        "D4": "Jakarta, 05 Juni 2026",
        "A5": "Attn",
        "D5": "Bapak Andi",
        "A6": "Email",
        "D6": "andi@kemala.co.id",
        "A8": "Deliv. Time",
        "D8": "3 working days",
        "A11": "No.",
        "B11": "Qty",
        "C11": "Unit",
        "D11": "Request",
        "E11": "IMPA",
        "F11": "Offer",
        "G11": "JUAL",
        "I11": "MODAL",
        "K11": "% profit",
        "M11": "Nama Asli barang",
        "N11": "Vendor",
        "O11": "Telp",
        "A13": 1,
        "B13": 2,
        "C13": "Pcs",
        "D13": "Shackle",
        "E13": 232101,
        "F13": "Shackle D 1 ton",
        "G13": 115000,
        "H13": 230000,
        "I13": 60000,
        "J13": 120000,
        "N13": "Toko B",
        "O13": "Toped",
        "A14": 2,
        "B14": 1,
        "C14": "Set",
        "D14": "Horn",
        "F14": "No Offer",
        "G14": 0,
        "H14": 0,
        "G20": "TOTAL",
        "H20": 230000,
    }
    printed = {
        "A7": "To", "D7": "PT Kemala Shipping/MBSS", "J7": "No", "L7": "Q-968208-O/GNS/VI/2026",
        "J8": "Your Ref No.", "L8": "PR08-9606-031",
        "J9": "Date", "L9": "Jakarta, 05 Juni 2026",
        "A12": "No.", "B12": "Qty", "C12": "Unit", "D12": "REQUEST", "H12": "IMPA",
        "I12": "O F F E R   D E S C R I P T I O N", "J12": "Unit Price", "L12": "Amount",
        "D13": "TB ASHLEIGH 06",
        "A15": 1, "B15": 2, "C15": "Pcs", "D15": "Shackle", "H15": 232101,
        "I15": "Shackle D 1 ton", "J15": 115000, "L15": 230000,
        "A16": 2, "B16": 1, "C16": "Set", "D16": "Horn", "I16": "No Offer", "J16": 0, "L16": 0,
        "J25": "Total", "L25": 230000,
        "J26": "Diskon 5%", "L26": -11500,
        "J27": "Sub Total", "L27": 218500,
        "J28": "DPP Nilai Lain", "L28": 200291.66666666666,
        "J29": "PPN 12%", "L29": 24035,
        "J30": "Grand Total", "L30": 242535,
        "A32": "DELIVERY PLACE", "E32": "Franco Jakarta",
        "A33": "DELIVERY TIME", "E33": "3 working days",
        "A34": "PAYMENT", "E34": "30 days",
        "A35": "VALIDITY", "E35": "3 days",
    }  # fmt: skip
    path = (
        tmp_path
        / "Quotation 2026 - Excel"
        / ("Q-968208-O (PR08-9606-031 - TB ASHLEIGH 06) - Kemala Shipping_MBSS.xlsx")
    )
    return save_workbook(path, {"DATA ENTRI": data, "PRINT HORIZONTAL": printed})


def test_build_workbook_record(tmp_path: Path) -> None:
    path = _workbook(tmp_path)
    records, excluded = build_workbook(path, tmp_path)
    assert excluded is None
    (q,) = records
    assert q["source"]["files"] == [
        "Quotation 2026 - Excel/Q-968208-O (PR08-9606-031 - TB ASHLEIGH 06) - Kemala Shipping_MBSS.xlsx"
    ]
    assert q["source"]["kind"] == "data_entry"
    assert q["number"]["original"] == "Q-968208-O/GNS/VI/2026"
    assert q["number"]["file"] == "Q-968208-O"
    assert q["date"] == "2026-06-05"
    assert q["client"] == {
        "name": "PT. Kemala Shipping",
        "raw": "PT Kemala Shipping/MBSS",
        "broker_note": "MBSS",
        "individual": False,
    }
    assert q["contact"] == {"name": "Bapak Andi", "email": "andi@kemala.co.id", "phone": None}
    assert (q["client_ref"], q["client_ref_source"]) == ("PR08-9606-031", "print")
    assert (q["vessel"], q["vessel_source"]) == ("TB ASHLEIGH 06", "print")
    assert q["terms"] == {
        "delivery_time": "3 working days",
        "delivery_place": "Franco Jakarta",
        "payment": "30 days",
        "payment_days": 30,
        "validity": "3 days",
        "validity_days": 3,
    }
    assert q["discount"] == {"kind": "pct", "pct": 5.0, "amount": 11500.0, "label": "Diskon 5%"}
    assert [(ln["line_no"], ln["available"]) for ln in q["lines"]] == [(1, True), (2, False)]
    assert q["lines"][0]["vendor"] == "Toko B" and q["lines"][0]["channel"] == "Toped"
    assert q["print_totals"]["grand"] == 242535
    assert q["computed_totals"]["grand"] == pytest.approx(242535, abs=0.01)
    assert q["flags"] == []


def test_build_workbook_reports_an_empty_second_sheet(tmp_path: Path) -> None:
    path = _workbook(tmp_path)
    copy = tmp_path / "Quotation 2026 - Excel" / "Q-968209 (PR08) - Kemala Shipping_MBSS.xlsx"
    import openpyxl

    wb = openpyxl.load_workbook(path)
    wb.create_sheet("DATA2")["A11"] = "No."
    wb["DATA2"]["B11"] = "Qty"
    wb["DATA2"]["D11"] = "Request"
    wb.save(copy)
    (q,) = build_workbook(copy, tmp_path)[0]
    assert "empty_sheet_dropped:DATA2" in q["flags"]


def _with_revision_sheet(tmp_path: Path, first: str, second: str) -> Path:
    import openpyxl

    path = _workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    sheet = wb.copy_worksheet(wb["PRINT HORIZONTAL"])
    sheet.title = "Print revisi"
    sheet["D15"] = sheet["I15"] = first
    sheet["D16"] = sheet["I16"] = second
    sheet["J15"] = 90000
    sheet["L15"] = 180000
    wb.save(path)
    return path


def test_build_workbook_binds_a_revision_sheet_to_its_data_sheet(tmp_path: Path) -> None:
    path = _with_revision_sheet(tmp_path, "Shackle", "Horn")
    data, revised = build_workbook(path, tmp_path)[0]
    assert revised["source"]["sheet"] == "Print revisi"
    assert revised["revision_of"] == data["id"]
    assert data["excluded_sheets"] == []


def test_build_workbook_drops_a_stale_revision_sheet(tmp_path: Path) -> None:
    # Q-655: "Print revisi" still holds Q-502's lines under Q-655's header.
    path = _with_revision_sheet(tmp_path, "Plasma Cutting CALDWELL CUT 100GE", "Headlamp LED")
    (data,) = build_workbook(path, tmp_path)[0]
    (excluded,) = data["excluded_sheets"]
    assert excluded["sheet"] == "Print revisi"
    assert excluded["reason"] == "stale revision sheet: no line in common with DATA ENTRI"
    assert excluded["lines"] == 2
    assert "stale_revision_sheet:Print revisi" in data["flags"]


def test_build_workbook_excludes_blank_template(tmp_path: Path) -> None:
    path = save_workbook(
        tmp_path / "Quotation 2025 - Excel" / "Q-952000 - Pelita Global Logistik.xlsx",
        {
            "DATA ENTRI": {
                "A2": "Customer", "D2": "PT. Pelita Global Logistik",
                "A3": "No", "D3": "Q-952000/GNS/I/2025",
                "A11": "No.", "B11": "Qty", "C11": "Unit", "D11": "D E S C R I P T I O N",
                "G11": "Harga Jual", "I11": "Modal", "A15": 1,
            },
            "PRINT": {},
        },
    )  # fmt: skip
    records, excluded = build_workbook(path, tmp_path)
    assert records == []
    assert excluded == "blank template (number 000, no lines)"


def test_build_workbook_excludes_non_quotation(tmp_path: Path) -> None:
    path = save_workbook(
        tmp_path / "Quotation 2024" / "Lumoso.xlsx",
        {"Sheet1": {"A1": "NO", "B1": "Request Items", "C1": "Quantity", "D1": "Modal"}},
    )
    records, excluded = build_workbook(path, tmp_path)
    assert records == []
    assert excluded == "not a quotation (no DATA ENTRI sheet and no quotation print sheet)"


def test_tax_setup_reads_the_print_block() -> None:
    flags: list[str] = []
    assert _tax_setup(None, 2024, flags) == (11.0, False, [])
    assert _tax_setup(None, 2026, flags) == (12.0, True, [])
    freight = PrintTotals(
        gross=1250000, discount=0, net=1250000, dpp=0, ppn=0, ppn_rate=12.0, ppn_label="PPN 12%"
    )
    assert _tax_setup(freight, 2025, flags) == (0.0, False, [])
    assert flags == ["ppn_zero_on_print"]
    shipping = PrintTotals(
        gross=1800000, discount=90000, net=2510000, dpp=2300833.3, ppn=276100, ppn_rate=12.0, ppn_label="PPN 12%",
        extras=[("Shipping to OBI", 800000.0, False)],
    )  # fmt: skip
    assert _tax_setup(shipping, 2025, [])[2] == [("Shipping to OBI", 800000.0, True)]


def test_tax_setup_without_ppn_row_and_with_withheld_pph() -> None:
    flags: list[str] = []
    services = PrintTotals(
        gross=25000000, discount=0, dpp=25000000, grand=24500000,
        extras=[("PPH 2%", -500000.0, False)],
    )  # fmt: skip
    rate, dpp, extras = _tax_setup(services, 2025, flags)
    assert (rate, dpp, extras, flags) == (
        0.0,
        False,
        [("PPH 2%", -500000.0, False)],
        ["no_ppn_on_print"],
    )
    totals = compute_totals(
        [{"line_total": 25000000.0, "in_total": True}],
        {"kind": "none", "pct": None, "amount": 0.0},
        rate,
        dpp,
        extras,
    )
    assert totals["grand"] == 24500000.0


def _rec_line(request: str, total: float) -> dict[str, object]:
    return {
        "row": 1,
        "no_raw": "1",
        "request": request,
        "offer": None,
        "sell": None,
        "sell_printed": None,
        "line_total": total,
        "in_total": True,
        "flags": [],
    }


def test_apply_print_line_set_excludes_lines_the_print_did_not_total() -> None:
    lines = [_rec_line("Relay A", 14000.0), _rec_line("Relay B", 14000.0), _rec_line("Seal", 160.0)]
    pr = PrintInfo(
        totals=PrintTotals(gross=14000.0),
        lines=[
            _line(request="Relay A", sell=7000.0, sell_amount=14000.0),
            _line(request="Seal", sell=32.0, sell_amount=None),
        ],
    )
    flags: list[str] = []
    _apply_print_line_set(lines, pr, flags)  # type: ignore[arg-type]
    assert [(ln["in_total"], ln["flags"]) for ln in lines] == [
        (True, []),
        (False, ["not_on_print"]),
        (False, ["no_amount_on_print"]),
    ]
    assert flags == ["lines_not_totalled_on_print"]


def test_apply_print_line_set_leaves_unexplained_gaps() -> None:
    lines = [_rec_line("Relay A", 14000.0), _rec_line("Relay B", 14000.0)]
    pr = PrintInfo(
        totals=PrintTotals(gross=20000.0),
        lines=[_line(request="Relay A", sell=7000.0, sell_amount=14000.0)],
    )
    flags: list[str] = []
    _apply_print_line_set(lines, pr, flags)  # type: ignore[arg-type]
    assert all(ln["in_total"] for ln in lines) and flags == []


def test_apply_print_line_set_matches_repeated_text_once_each() -> None:
    first = {**_rec_line("Hawser PP 8-strand", 65000.0), "sell": 32500.0, "sell_printed": 32500.0}
    second = {**_rec_line("Hawser PP 8-strand", 90000.0), "sell": 45000.0, "sell_printed": 45000.0}
    pr = PrintInfo(
        totals=PrintTotals(gross=65000.0),
        lines=[_line(request="Hawser PP 8-strand", sell=32500.0, sell_amount=65000.0)],
    )
    _apply_print_line_set([first, second], pr, [])  # type: ignore[arg-type]
    assert (first["in_total"], second["in_total"]) == (True, False)
    assert second["flags"] == ["not_on_print"]


def test_apply_print_line_set_keeps_a_price_list_with_a_zero_total() -> None:
    # Q-965127: option prices with no amounts, so PRINT totals Rp 0.
    lines = [_rec_line("Compressor A", 1500000.0), _rec_line("Compressor B", 1400000.0)]
    pr = PrintInfo(
        totals=PrintTotals(gross=0.0),
        lines=[
            _line(request="Compressor A", sell=1500000.0, sell_amount=None),
            _line(request="Compressor B", sell=1400000.0, sell_amount=None),
        ],
    )
    flags: list[str] = []
    _apply_print_line_set(lines, pr, flags)  # type: ignore[arg-type]
    assert all(ln["in_total"] and not ln["flags"] for ln in lines)
    assert flags == ["print_price_list"]


def test_apply_print_line_set_flags_a_zero_total_with_amounts() -> None:
    lines = [_rec_line("Seal", 312000.0)]
    pr = PrintInfo(totals=PrintTotals(gross=0.0), lines=[_line(request="Other seal", sell=0.0)])
    flags: list[str] = []
    _apply_print_line_set(lines, pr, flags)  # type: ignore[arg-type]
    assert lines[0]["in_total"] and flags == ["print_total_zero"]


def test_apply_print_no_offer_uses_the_printed_offer() -> None:
    locks = normalize_line(_line(request="Mortise lock", offer=None, sell=0.0), has_amount=True)
    fan = normalize_line(_line(request="Blower", offer="Fan 220V", sell=5.0), has_amount=True)
    pr = PrintInfo(
        lines=[
            _line(request="Mortise lock", offer="No Offer", sell=0.0, sell_amount=0.0),
            _line(request="Blower", offer="Fan 220V", sell=5.0, sell_amount=10.0),
        ]
    )
    _apply_print_no_offer([locks, fan], pr)
    assert (locks["available"], locks["flags"]) == (False, ["no_offer_on_print"])
    assert fan["available"] is True
