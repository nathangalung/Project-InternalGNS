"""Readers for client PO and invoice source files.

Two sources are read mechanically:

- Pelita Global Logistik portal orders, from `pdftotext -layout` text.
- GNS DO and invoice workbooks, from the cached cell values plus the cell
  formulas (the formulas say which cell is a quantity, a price or a total).

Everything returned is as printed; judgement lives in build_docs.py.
"""

from __future__ import annotations

import re
import subprocess
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

CENT = Decimal("0.01")

MONTHS: dict[str, int] = {
    "jan": 1,
    "januari": 1,
    "january": 1,
    "feb": 2,
    "februari": 2,
    "february": 2,
    "mar": 3,
    "maret": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "mei": 5,
    "may": 5,
    "jun": 6,
    "juni": 6,
    "june": 6,
    "jul": 7,
    "juli": 7,
    "july": 7,
    "agu": 8,
    "agt": 8,
    "agustus": 8,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "okt": 10,
    "oktober": 10,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "des": 12,
    "desember": 12,
    "dec": 12,
    "december": 12,
}

INV_NO = re.compile(r"\b(\d{3,5}[A-Z]?/INV-GNS/[IVX]+/\d{4})")
DO_NO = re.compile(r"\b(\d{3,5}[A-Z]?/DO-GNS/[IVX]+/\d{4})")
PO_NO = re.compile(
    r"(JKT-PO/[0-9./A-Z]+"
    r"|[VO]-\d{2}-\d{4}-\d{3}-[DE]/\d{2}/\d{2}"
    r"|(?:PO|WO)-(?:TCP|SML)/[IVX]+/\d{4}-\d{5}"
    r"|ID-\d{9}"
    r"|8404/O-\d{4}/P\d{3})"
)


def money(value: object) -> Decimal:
    """Round a cell or text amount to the sen."""
    if isinstance(value, Decimal):
        return value.quantize(CENT, ROUND_HALF_UP)
    if isinstance(value, (int, float)):
        return Decimal(repr(value)).quantize(CENT, ROUND_HALF_UP)
    if isinstance(value, str):
        return Decimal(value.replace(",", "")).quantize(CENT, ROUND_HALF_UP)
    raise TypeError(f"not an amount: {value!r}")


def plain(value: Decimal) -> Decimal:
    """Drop trailing zeros so 3.00 compares and prints as 3."""
    return value.quantize(Decimal(1)) if value == value.to_integral() else value.normalize()


# Dates
def parse_id_date(text: str) -> date | None:
    """Find one Indonesian or English date in free text."""
    s = text.strip()
    m = re.search(r"(\d{4})/(\d{1,2})/(\d{1,2})", s)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    # The year may follow the month unspaced: "29 December2025"
    m = re.search(r"(\d{1,2})[\s-]+([A-Za-z]+)[\s-]*(\d{4})", s)
    if m and m[2].lower() in MONTHS:
        return date(int(m[3]), MONTHS[m[2].lower()], int(m[1]))
    m = re.search(r"([A-Za-z]+)\s+(\d{1,2}),\s*(\d{4})", s)
    if m and m[1].lower() in MONTHS:
        return date(int(m[3]), MONTHS[m[1].lower()], int(m[2]))
    return None


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return parse_id_date(value)
    return None


# Pelita portal orders
@dataclass
class PortalLine:
    position: int
    description: str
    code: str
    qty: Decimal
    unit: str
    unit_price: Decimal
    amount: Decimal
    remark: str | None = None
    supplier_remark: str | None = None


@dataclass
class PortalOrder:
    ref_no: str
    order_date: date | None
    vessel: str | None
    order_no: str | None
    offer_no: str | None
    city: str | None
    price: Decimal
    rebate_pct: Decimal
    rebate: Decimal
    charges: list[tuple[str, Decimal]]
    sub_total: Decimal
    vat_pct: Decimal
    vat: Decimal
    total: Decimal
    payment_terms: str | None = None
    rfq_date: date | None = None
    lines: list[PortalLine] = field(default_factory=list)


AMT = r"(-?[\d,]+\.\d{2})\s+IDR"
LINE_RE = re.compile(
    r"^(?P<pos>\d{3})\s{2,}(?P<mid>.+?)\s+(?P<qty>\d[\d.,]*)\s+(?P<unit>\S+)\s+"
    r"(?P<price>[\d,]+\.\d{4})(?:\s+(?P<pct>[\d.]+))?\s+(?P<amt>[\d,]+\.\d{4})\s*$"
)
SUMMARY_SKIP = ("Price", "Rebate", "Sub Total", "VAT", "Total")
REMARK_LABELS = ("Purchaser:", "Supplier:")


def _field(text: str, label: str) -> str | None:
    m = re.search(rf"^{re.escape(label)}[ \t]{{1,30}}(\S.*?)(?:\s{{2,}}|$)", text, re.M)
    return m[1].strip() if m else None


def _summary(text: str, label: str) -> Decimal:
    m = re.search(rf"{re.escape(label)}.*?{AMT}", text)
    if not m:
        raise ValueError(f"price summary has no {label!r}")
    return money(m[1])


def parse_pelita_order(text: str) -> PortalOrder:
    """Parse one portal order rendered by `pdftotext -layout`."""
    ref = re.search(r"Reference No\. / Case ID\s+(\S+)", text)
    if not ref:
        raise ValueError("not a portal order: no Reference No.")
    od = re.search(r"Order Date\s+([A-Z][a-z]{2} \d{1,2}, \d{4})", text)
    rebate = re.search(r"Rebate \[%\]:\s+([\d.]+)%\s+" + AMT, text)
    vat = re.search(r"VAT \[%\]:\s+([\d.]+)%\s+" + AMT, text)
    city = re.search(r"Country, City:[ \t]+(\S.*?)\s*$", text, re.M)
    terms = re.search(r"Payment Terms:[ \t]+(\S.*?)\s*$", text, re.M)
    rfq = _field(text, "RFQ Date:")
    charges: list[tuple[str, Decimal]] = []
    block = text[text.index("Price Summary:") : text.index("Sub Total:")]
    for m in re.finditer(r"^\s+([A-Za-z][A-Za-z &]+?):?\s+" + AMT + r"\s*$", block, re.M):
        label, amount = m[1].strip(), money(m[2])
        if label.startswith(SUMMARY_SKIP):
            continue
        if amount or label not in ("Transport", "Packing", "Insurance"):
            charges.append((label, amount))
    totals = re.findall(r"^\s+Total:\s+" + AMT, text, re.M)
    if not totals:
        raise ValueError("price summary has no Total")
    return PortalOrder(
        ref_no=ref[1],
        order_date=parse_id_date(od[1]) if od else None,
        vessel=_field(text, "Name:"),
        order_no=_field(text, "Order No.:"),
        offer_no=_field(text, "Offer No.:"),
        city=city[1].strip() if city else None,
        price=_summary(text, "Price:"),
        rebate_pct=Decimal(rebate[1]) if rebate else Decimal(0),
        rebate=money(rebate[2]) if rebate else Decimal(0),
        charges=charges,
        sub_total=_summary(text, "Sub Total:"),
        vat_pct=plain(Decimal(vat[1])) if vat else Decimal(0),
        vat=money(vat[2]) if vat else Decimal(0),
        total=money(totals[-1]),
        payment_terms=terms[1] if terms else None,
        rfq_date=parse_id_date(rfq) if rfq else None,
        lines=_portal_lines(text),
    )


def _portal_lines(text: str) -> list[PortalLine]:
    rows = text.splitlines()
    out: list[PortalLine] = []
    i = 0
    while i < len(rows):
        m = LINE_RE.match(rows[i])
        if not m:
            i += 1
            continue
        parts = [p for p in re.split(r"\s{2,}", m["mid"].strip()) if p]
        if len(parts) > 1:
            desc, code = " ".join(parts[:-1]), parts[-1]
        else:
            desc, _, code = parts[0].rpartition(" ")
        code_col = rows[i].index(code, m.start("mid") + len(desc))
        desc_lines = [desc.strip()]
        j = i + 1
        while j < len(rows):
            left = rows[j][:code_col].strip()
            if not left or left.startswith(REMARK_LABELS) or LINE_RE.match(rows[j]):
                break
            desc_lines.append(left)
            j += 1
        code = code.strip()
        # A long name pushes the group number (code[:2]) into it
        group = f" {code[:2]}"
        if len(desc_lines) > 1 and code.isdigit() and desc_lines[1].endswith(group):
            desc_lines[1] = desc_lines[1][: -len(group)].rstrip()
        end = next((k for k in range(j, len(rows)) if LINE_RE.match(rows[k])), len(rows))
        remarks = _remarks(rows[j:end])
        out.append(
            PortalLine(
                position=int(m["pos"]),
                description="\n".join(desc_lines),
                code=code,
                qty=plain(Decimal(m["qty"].replace(",", ""))),
                unit=m["unit"],
                unit_price=plain(money(m["price"])),
                amount=plain(money(m["amt"])),
                remark=remarks.get("Purchaser:"),
                supplier_remark=remarks.get("Supplier:"),
            )
        )
        i = j
    return out


def _remarks(rows: list[str]) -> dict[str, str]:
    """Purchaser and Supplier notes, with their indented continuation rows."""
    found: dict[str, str] = {}
    k = 0
    while k < len(rows):
        label = next((x for x in REMARK_LABELS if rows[k].strip().startswith(x)), None)
        if label is None:
            k += 1
            continue
        start = rows[k].index(label) + len(label)
        first = rows[k][start:]
        col = start + len(first) - len(first.lstrip())
        parts = [first.strip()]
        k += 1
        while k < len(rows) and rows[k].strip() and not rows[k][:col].strip():
            parts.append(rows[k].strip())
            k += 1
        found[label] = "\n".join(parts)
    return found


def pdf_text(path: Path) -> str:
    """Return the layout text of a PDF."""
    res = subprocess.run(
        ["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True, check=True
    )
    return res.stdout


# Invoice workbooks
@dataclass
class InvoiceLine:
    row: int
    description: str | None
    qty: Decimal
    unit: str | None
    unit_price: Decimal
    amount: Decimal


@dataclass
class Buyer:
    """The billed party as the invoice prints it."""

    name: str
    address: str | None
    npwp: str | None


@dataclass
class InvoiceDoc:
    invoice_numbers: list[str]
    do_numbers: list[str]
    invoice_date: date | None
    due_date: date | None
    po_number: str | None
    po_date: date | None
    do_date: date | None
    client: str | None
    lines: list[InvoiceLine]
    subtotal: Decimal
    discount: Decimal
    ppn: Decimal
    total: Decimal
    ppn_rate: Decimal
    buyer: Buyer | None = None


PRODUCT = re.compile(r"^=\(?([A-Z]+)(\d+)\*([A-Z]+)(\d+)\)?$")
DO_REF = re.compile(r"^=(?:'?DO'?)!([A-Z]+)(\d+)$")


def _sheet(wb: Workbook, name: str) -> Worksheet | None:
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name:
            return ws
    return None


def _strings(ws: Worksheet, pattern: re.Pattern[str]) -> list[tuple[str, str]]:
    return [
        (c.coordinate, m[1])
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str)
        for m in pattern.finditer(c.value)
    ]


def _uniq(values: list[str]) -> list[str]:
    return list(dict.fromkeys(values))


def _col(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def _num(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def read_invoice(values: Workbook, formulas: Workbook) -> InvoiceDoc:
    """Read one DO and invoice workbook as printed."""
    inv_v = _sheet(values, "invoice")
    inv_f = _sheet(formulas, "invoice")
    if inv_v is None or inv_f is None:
        raise ValueError("workbook has no Invoice sheet")
    do_v = _sheet(values, "do")

    inv_hits = _strings(inv_v, INV_NO)
    if not inv_hits:
        raise ValueError("Invoice sheet prints no invoice number")
    do_hits = _strings(inv_v, DO_NO) or (_strings(do_v, DO_NO) if do_v is not None else [])
    inv_cell = inv_v[inv_hits[0][0]]

    lines = _invoice_lines(inv_v, inv_f, do_v)
    sub, disc, ppn, total, ppn_cell = _totals(inv_v, lines)
    ppn_form = inv_f[ppn_cell].value
    zero_vat = ppn == 0 or (isinstance(ppn_form, str) and ppn_form.replace(" ", "").endswith("*0"))

    # A PO number inside a line's text is not the PO reference
    line_rows = {ln.row for ln in lines}
    po_hits = [h for h in _strings(inv_v, PO_NO) if inv_v[h[0]].row not in line_rows]
    po_number = po_hits[0][1] if po_hits else None
    po_date = None
    if po_hits:
        pc = inv_v[po_hits[0][0]]
        po_date = as_date(inv_v.cell(pc.row + 1, pc.column).value)

    inv_date = _date_above(inv_v, inv_cell.row, inv_cell.column)
    due = _due_date(inv_v, inv_f, inv_date)
    return InvoiceDoc(
        invoice_numbers=_uniq([h[1] for h in inv_hits]),
        do_numbers=_uniq([h[1] for h in do_hits]),
        invoice_date=inv_date,
        due_date=due,
        po_number=po_number,
        po_date=po_date,
        do_date=_do_date(do_v),
        client=_client(inv_v),
        lines=lines,
        subtotal=sub,
        discount=abs(disc),
        ppn=ppn,
        total=total,
        ppn_rate=Decimal(0) if zero_vat else Decimal(11),
        buyer=buyer_block(inv_v) or (buyer_block(do_v) if do_v is not None else None),
    )


def _date_above(ws: Worksheet, row: int, col: int) -> date | None:
    for r in range(row - 1, 0, -1):
        d = as_date(ws.cell(r, col).value)
        if d:
            return d
    return None


def _due_date(inv_v: Worksheet, inv_f: Worksheet, inv_date: date | None) -> date | None:
    for row in inv_f.iter_rows():
        for c in row:
            if isinstance(c.value, str) and re.match(r"^=[A-Z]+\d+\+\d+$", c.value):
                return as_date(inv_v[c.coordinate].value)
    for row in inv_v.iter_rows(max_row=16):
        for c in row:
            d = as_date(c.value) if isinstance(c.value, str) else None
            if d and d != inv_date:
                return d
    return None


def _do_date(do_v: Worksheet | None) -> date | None:
    if do_v is None:
        return None
    below_po: set[tuple[int, int]] = set()
    for coord, _ in _strings(do_v, PO_NO):
        c = do_v[coord]
        below_po.add((c.row + 1, c.column))
    for row in do_v.iter_rows(max_row=16):
        for c in row:
            if (c.row, c.column) in below_po:
                continue
            d = as_date(c.value)
            if d:
                return d
    return None


# A phone line or a contact person ends the address.
_CONTACT_LINE = re.compile(
    r"^[\d\s()+/.-]{7,}$|^(bp|bpk|bapak|ibu|pak|mr|mrs|ms)\b|\b0\d{8,}|telp|phone|fax",
    re.IGNORECASE,
)
_NPWP_LINE = re.compile(r"NPWP\s*:?\s*([\d.\-\s]{15,})", re.IGNORECASE)


def buyer_block(ws: Worksheet) -> Buyer | None:
    """The buyer block under the billed company: name, address lines, NPWP.

    The company name opens the block; the lines below it in the same
    column are the address until a phone or contact line. A second company
    after a slash or a line break is a brokered party, not the buyer.
    """
    rows = list(ws.iter_rows(max_row=20))
    for i, row in enumerate(rows):
        for c in row:
            v = c.value
            if not (
                isinstance(v, str)
                and re.match(r"\s*PT[ .]", v)
                and "GLOBAL NIAGA" not in v.upper()
                and "PEMBAYARAN" not in v.upper()
            ):
                continue
            name = re.split(r"\n|/\s*PT\b", v.strip())[0].strip()
            address: list[str] = []
            npwp = None
            for below in rows[i + 1 : i + 6]:
                cell = next((x for x in below if x.column == c.column), None)
                text = (
                    cell.value.strip() if cell is not None and isinstance(cell.value, str) else ""
                )
                if not text:
                    break
                tax = _NPWP_LINE.search(text)
                if tax:
                    npwp = re.sub(r"\D", "", tax.group(1))
                    continue
                if _CONTACT_LINE.search(text):
                    break
                address.append(text.rstrip(" ,"))
            return Buyer(name, ", ".join(address) or None, npwp)
    return None


def read_do_units(path: Path) -> list[str | None]:
    """The unit of each item row on a workbook's DO sheet, in order."""
    do_v = _sheet(load_workbook(path, data_only=True), "do")
    if do_v is None:
        return []
    units: list[str | None] = []
    for row in do_v.iter_rows():
        cells = [c.value for c in row[1:5]]
        if len(cells) == 4 and _num(cells[0]) and _num(cells[1]) and isinstance(cells[3], str):
            units.append(
                cells[2].strip() if isinstance(cells[2], str) and cells[2].strip() else None
            )
    return units


def _client(ws: Worksheet) -> str | None:
    for row in ws.iter_rows(max_row=16):
        for c in row:
            v = c.value
            if (
                isinstance(v, str)
                and re.match(r"\s*PT[ .]", v)
                and "GLOBAL NIAGA" not in v.upper()
                and "PEMBAYARAN" not in v.upper()
            ):
                return v.strip()
    return None


def _invoice_lines(inv_v: Worksheet, inv_f: Worksheet, do_v: Worksheet | None) -> list[InvoiceLine]:
    lines: list[InvoiceLine] = []
    for row in inv_f.iter_rows():
        for c in row:
            f = c.value
            m = PRODUCT.match(f.replace(" ", "")) if isinstance(f, str) else None
            if not m:
                continue
            a = (_col(m[1]), int(m[2]))
            b = (_col(m[3]), int(m[4]))
            if c.row not in (a[1], b[1]):
                continue
            qty_ref, price_ref = sorted((a, b))
            amount = inv_v[c.coordinate].value
            price = inv_v.cell(price_ref[1], price_ref[0]).value
            if not _num(amount) or not _num(price):
                continue
            desc = _description(inv_v, c.row, qty_ref[0], price_ref[0])
            if amount == 0 and desc is None:
                continue
            qty_cell = inv_v.cell(c.row, qty_ref[0]).value
            qty = Decimal(repr(qty_cell)) if _num(qty_cell) else money(amount) / money(price)
            lines.append(
                InvoiceLine(
                    row=c.row,
                    description=desc,
                    qty=plain(qty),
                    unit=_unit(inv_f, do_v, c.row, qty_ref[0]),
                    unit_price=plain(money(price)),
                    amount=plain(money(amount)),
                )
            )
    if lines:
        return lines
    # Hard-coded rows: no product formula
    for row in inv_v.iter_rows():
        nums = [(c.column, c.value) for c in row if _num(c.value)]
        if len(nums) >= 3:
            (qc, q), (_, p), (_, a) = nums[0], nums[-2], nums[-1]
            if q and abs(q * p - a) < 0.005:
                lines.append(
                    InvoiceLine(
                        row=row[0].row,
                        description=_description(inv_v, row[0].row, qc, nums[-2][0]),
                        qty=plain(Decimal(repr(q))),
                        unit=None,
                        unit_price=plain(money(p)),
                        amount=plain(money(a)),
                    )
                )
    return lines


def _description(ws: Worksheet, row: int, lo: int, hi: int) -> str | None:
    for col in range(lo + 1, hi):
        v = ws.cell(row, col).value
        if isinstance(v, str) and v.strip():
            return "\n".join(part.strip() for part in v.strip().splitlines() if part.strip())
    return None


def _unit(inv_f: Worksheet, do_v: Worksheet | None, row: int, qty_col: int) -> str | None:
    f = inv_f.cell(row, qty_col).value
    m = DO_REF.match(f) if isinstance(f, str) else None
    if m and do_v is not None:
        v = do_v.cell(int(m[2]), _col(m[1]) + 1).value
        if isinstance(v, str) and v.strip():
            return v.strip()
    return None


def _totals(
    inv_v: Worksheet, lines: list[InvoiceLine]
) -> tuple[Decimal, Decimal, Decimal, Decimal, str]:
    """Find the subtotal, discount, PPN and total block below the lines."""
    line_sum = sum((ln.amount for ln in lines), Decimal(0))
    last = max((ln.row for ln in lines), default=0)
    for col in range(1, inv_v.max_column + 1):
        for r in range(last + 1, inv_v.max_row - 2):
            vals = [inv_v.cell(r + k, col).value for k in range(4)]
            if not all(_num(v) for v in vals):
                continue
            sub, disc, ppn, tot = (money(v) for v in vals)
            if sub != line_sum:
                continue
            if tot in (sub + disc + ppn, sub - disc + ppn):
                return sub, disc, ppn, tot, inv_v.cell(r + 2, col).coordinate
    raise ValueError(f"no totals block sums the lines ({line_sum})")


def read_invoice_file(path: Path) -> InvoiceDoc:
    """Load a workbook twice (values and formulas) and read it."""
    return read_invoice(load_workbook(path, data_only=True), load_workbook(path, data_only=False))


def read_do_sheet(values: Workbook) -> tuple[list[str], date | None]:
    """The DO numbers and date printed on a workbook's DO sheet."""
    do_v = _sheet(values, "do")
    if do_v is None:
        raise ValueError("workbook has no DO sheet")
    return _uniq([h[1] for h in _strings(do_v, DO_NO)]), _do_date(do_v)


def read_do_file(path: Path) -> tuple[list[str], date | None]:
    return read_do_sheet(load_workbook(path, data_only=True))


# Quotation workbooks
QUOTE_NO = re.compile(r"(Q-[0-9A-Z-]+(?:/\s*\d*GNS)?/[IVX]+/\d{4})")


@dataclass
class QuoteRow:
    row: int
    qty: Decimal | None
    text: str | None
    price: Decimal | None
    numbers: list[Decimal]


@dataclass
class Quote:
    sheet: str
    printed_number: str | None
    printed_date: date | None
    rows: list[QuoteRow]
    price_column: bool

    def rows_priced(self, price: Decimal) -> list[QuoteRow]:
        """Rows whose sell price is this amount (any number without a sell column)."""
        if self.price_column:
            return [r for r in self.rows if r.price == price]
        return [r for r in self.rows if price in r.numbers]


def _number(value: object) -> Decimal | None:
    if _num(value):
        return money(value)
    if isinstance(value, str) and re.fullmatch(r"\s*-?\d+(?:\.\d+)?\s*", value):
        return money(value.strip())
    return None


def read_quote(wb: Workbook) -> Quote:
    """Read the DATA ENTRI sheet of one quotation workbook."""
    ws = next((s for s in wb.worksheets if "DATA" in s.title.upper()), wb.worksheets[0])
    printed = None
    when = None
    for row in ws.iter_rows(max_row=8):
        for c in row:
            if isinstance(c.value, str):
                m = QUOTE_NO.search(c.value)
                if m and printed is None:
                    printed = m[1]
                elif when is None and re.search(r"\d{4}", c.value):
                    when = parse_id_date(c.value)
            elif when is None and isinstance(c.value, datetime):
                when = c.value.date()
    header = 0
    cols: dict[str, int] = {}
    for row in ws.iter_rows(max_row=30):
        for c in row:
            if isinstance(c.value, str) and c.value.strip().lower() == "qty":
                header = c.row
        if header:
            for c in row:
                if isinstance(c.value, str):
                    key = c.value.strip().lower().replace(" ", "")
                    if key == "qty":
                        cols["qty"] = c.column
                    elif key.startswith("offer"):
                        cols["offer"] = c.column
                    elif key.startswith(("request", "description", "descr")):
                        cols["request"] = c.column
                    elif key.startswith(("hargajual", "jual", "sell")):
                        cols["price"] = c.column
            break
    rows: list[QuoteRow] = []
    if header:
        for row in ws.iter_rows(min_row=header + 1):
            texts = [c.value.strip() for c in row if isinstance(c.value, str) and c.value.strip()]
            if any(t.upper() == "TOTAL" for t in texts):
                break
            qty = _number(ws.cell(row[0].row, cols["qty"]).value) if "qty" in cols else None
            numbers = [
                n
                for c in row
                if c.column != cols.get("qty") and (n := _number(c.value)) is not None
            ]
            if qty is None and not numbers:
                continue
            text = None
            for key in ("offer", "request"):
                v = ws.cell(row[0].row, cols[key]).value if key in cols else None
                if isinstance(v, str) and v.strip():
                    text = v.strip()
                    break
            price = _number(ws.cell(row[0].row, cols["price"]).value) if "price" in cols else None
            rows.append(
                QuoteRow(
                    row=row[0].row,
                    qty=plain(qty) if qty is not None else None,
                    text=text,
                    price=plain(price) if price is not None else None,
                    numbers=[plain(n) for n in numbers],
                )
            )
    if when is None:
        raise ValueError(f"{ws.title}: no quotation date in the header")
    return Quote(
        sheet=ws.title,
        printed_number=printed,
        printed_date=when,
        rows=rows,
        price_column="price" in cols,
    )


def read_quote_file(path: Path) -> Quote:
    return read_quote(load_workbook(path, data_only=True))


def workbook_cells(path: Path) -> dict[tuple[str, str], object]:
    """Every non-empty cached value, keyed by sheet and cell."""
    wb = load_workbook(path, data_only=True)
    return {
        (ws.title, c.coordinate): c.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if c.value not in (None, "")
    }
